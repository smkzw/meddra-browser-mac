from __future__ import annotations

import csv
import io
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

from .meddra_data import LEVEL_LABELS, fuzzy_candidate, fuzzy_score, fuzzy_threshold, normalize_text, token_key


# Data listing columns that usually need MedDRA coding (label fragment → coding type).
# Priority 1: explicit field OIDs in parentheses (most reliable on EDC listings).
OID_CODING_MAP: dict[str, str] = {
    "AETERM": "AE",
    "MHTERM": "MH",
    "CMINDC": "CM",
    "AHDESC": "AH",
    "ALRTERM": "ALR",
    "DSDECOD": "DS",
}

# Priority 2: sheet-specific Chinese labels. Must appear as the primary label
# (not inside 是非问句 / 编号 / 备注).
CODING_COLUMN_HINTS: dict[str, tuple[str, ...]] = {
    "AE": ("不良事件名称", "不良事件描述"),
    "MH": ("疾病名称", "既往病史名称", "现病史名称"),
    "CM": ("用药原因", "适应症"),
    "AH": ("过敏史详述", "过敏描述"),
    "ALR": ("项目名称",),
    "DS": ("提前退出原因",),
}

# Headers that look similar but are never coding targets.
SKIP_HEADER_PATTERNS = (
    re.compile(r"^是否"),
    re.compile(r"编号$"),
    re.compile(r"^(项目编号|表单编号|受试者编号)$"),
    re.compile(r"(日期|时间|备注|说明|其他)"),
    re.compile(r"(剂量|频率|途径|严重|程度|转归|关系|措施|标准|状态)"),
)

# Header cells that identify the listing metadata columns.
META_HEADERS = {
    "项目编号",
    "表单编号",
    "受试者编号",
    "姓名缩写",
    "受试者状态",
    "试验中心编号",
    "试验中心名称",
    "数据节",
    "Instance顺序号",
    "数据块",
    "Block顺序号",
    "数据页",
    "最后修改时间",
    "行号",
}

SESSION_META_HEADERS = [
    "项目编号",
    "表单编号",
    "受试者编号",
    "受试者状态",
    "试验中心编号",
    "试验中心名称",
    "数据节",
    "行号",
]


@dataclass
class CodingTarget:
    sheet: str
    column: str
    coding_type: str
    row_count: int
    non_empty_count: int

    def as_dict(self) -> dict[str, Any]:
        return {
            "sheet": self.sheet,
            "column": self.column,
            "coding_type": self.coding_type,
            "row_count": self.row_count,
            "non_empty_count": self.non_empty_count,
        }


@dataclass
class UniqueTerm:
    term: str
    coding_type: str
    count: int
    occurrences: list[dict[str, Any]] = field(default_factory=list)

    def as_dict(self) -> dict[str, Any]:
        return {
            "term": self.term,
            "coding_type": self.coding_type,
            "count": self.count,
            "occurrences": self.occurrences,
        }


def _header_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _match_coding_type(sheet: str, header: str) -> str | None:
    sheet_upper = (sheet or "").strip().upper()
    header = (header or "").strip()
    if not header or header in META_HEADERS:
        return None
    for pattern in SKIP_HEADER_PATTERNS:
        if pattern.search(header):
            return None

    # Strongest signal: EDC field OID in trailing parentheses.
    oid_match = re.search(r"\(([A-Za-z0-9_]+)\)\s*$", header)
    if oid_match:
        oid = oid_match.group(1).upper()
        if oid in OID_CODING_MAP:
            return OID_CODING_MAP[oid]

    # Sheet-scoped Chinese label (exact fragment at start or as the bare label).
    for coding_type, fragments in CODING_COLUMN_HINTS.items():
        if sheet_upper and (sheet_upper == coding_type or sheet_upper.startswith(coding_type)):
            for fragment in fragments:
                if header == fragment or header.startswith(fragment):
                    return coding_type
    return None


def _row_meta(row: dict[str, Any]) -> dict[str, Any]:
    meta: dict[str, Any] = {}
    for key in SESSION_META_HEADERS:
        if key in row:
            value = row[key]
            meta[key] = "" if value is None else str(value)
    return meta


def _parse_rows_from_matrix(headers: list[str], matrix: list[list[Any]], sheet: str) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for index, values in enumerate(matrix):
        row: dict[str, Any] = {"_sheet": sheet, "_excel_row": index + 2}
        for col_index, header in enumerate(headers):
            if not header:
                continue
            row[header] = values[col_index] if col_index < len(values) else ""
        rows.append(row)
    return rows


def _detect_targets(sheet: str, headers: list[str]) -> list[tuple[str, str]]:
    targets: list[tuple[str, str]] = []
    for header in headers:
        coding_type = _match_coding_type(sheet, header)
        if coding_type:
            targets.append((header, coding_type))
    return targets


def parse_data_listing_excel(content: bytes, filename: str = "") -> dict[str, Any]:
    """Parse a multi-sheet data listing workbook (or CSV) into coding targets + unique terms."""
    from openpyxl import load_workbook

    if filename.lower().endswith(".csv") or (b"\x00" not in content[:200] and b"," in content[:200] and b"PK" not in content[:4]):
        return parse_data_listing_csv(content, filename or "listing.csv")

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheets_info: list[dict[str, Any]] = []
        targets: list[CodingTarget] = []
        unique_map: dict[tuple[str, str], UniqueTerm] = {}
        total_rows = 0

        for sheet_name in workbook.sheetnames:
            if sheet_name.upper() == "TOC":
                continue
            ws = workbook[sheet_name]
            rows_iter = ws.iter_rows(values_only=True)
            try:
                header_row = next(rows_iter)
            except StopIteration:
                continue
            headers = [_header_text(cell) for cell in header_row]
            if not any(headers):
                continue

            sheet_rows: list[list[Any]] = []
            for values in rows_iter:
                if values is None:
                    continue
                # Skip fully empty rows.
                if all(cell is None or str(cell).strip() == "" for cell in values):
                    continue
                sheet_rows.append(list(values))
            total_rows += len(sheet_rows)
            rows = _parse_rows_from_matrix(headers, sheet_rows, sheet_name)
            sheet_targets = _detect_targets(sheet_name, headers)
            sheets_info.append(
                {
                    "sheet": sheet_name,
                    "headers": headers,
                    "row_count": len(rows),
                    "coding_columns": [{"column": col, "coding_type": ctype} for col, ctype in sheet_targets],
                }
            )

            for column, coding_type in sheet_targets:
                non_empty = 0
                for row in rows:
                    value = row.get(column)
                    if value is None:
                        continue
                    term = str(value).strip()
                    if not term:
                        continue
                    non_empty += 1
                    key = (term, coding_type)
                    if key not in unique_map:
                        unique_map[key] = UniqueTerm(term=term, coding_type=coding_type, count=0)
                    unique_map[key].count += 1
                    unique_map[key].occurrences.append(
                        {
                            "sheet": sheet_name,
                            "column": column,
                            **_row_meta(row),
                            "excel_row": row.get("_excel_row"),
                        }
                    )
                targets.append(
                    CodingTarget(
                        sheet=sheet_name,
                        column=column,
                        coding_type=coding_type,
                        row_count=len(rows),
                        non_empty_count=non_empty,
                    )
                )

        unique_terms = sorted(
            unique_map.values(),
            key=lambda item: (-item.count, item.coding_type, item.term),
        )
        return {
            "filename": filename,
            "sheets": sheets_info,
            "coding_targets": [t.as_dict() for t in targets],
            "unique_terms": [t.as_dict() for t in unique_terms],
            "stats": {
                "sheet_count": len(sheets_info),
                "data_rows": total_rows,
                "target_count": len(targets),
                "unique_term_count": len(unique_terms),
            },
        }
    finally:
        workbook.close()


def parse_data_listing_csv(content: bytes, filename: str = "listing.csv") -> dict[str, Any]:
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    matrix = [list(row) for row in reader]
    if not matrix:
        return {
            "filename": filename,
            "sheets": [],
            "coding_targets": [],
            "unique_terms": [],
            "stats": {"sheet_count": 0, "data_rows": 0, "target_count": 0, "unique_term_count": 0},
        }
    headers = [_header_text(cell) for cell in matrix[0]]
    body = [list(row) for row in matrix[1:] if any(str(c).strip() for c in row)]
    # Treat CSV as a single sheet named after the file stem.
    sheet_name = Path(filename).stem[:32] or "CSV"
    rows = _parse_rows_from_matrix(headers, body, sheet_name)
    targets: list[CodingTarget] = []
    unique_map: dict[tuple[str, str], UniqueTerm] = {}
    sheet_targets = _detect_targets(sheet_name, headers)

    # If no sheet-based match, try header-only detection.
    if not sheet_targets:
        sheet_targets = _detect_targets("", headers)

    for column, coding_type in sheet_targets:
        non_empty = 0
        for row in rows:
            value = row.get(column)
            if value is None:
                continue
            term = str(value).strip()
            if not term:
                continue
            non_empty += 1
            key = (term, coding_type)
            if key not in unique_map:
                unique_map[key] = UniqueTerm(term=term, coding_type=coding_type, count=0)
            unique_map[key].count += 1
            unique_map[key].occurrences.append(
                {
                    "sheet": sheet_name,
                    "column": column,
                    **_row_meta(row),
                    "excel_row": row.get("_excel_row"),
                }
            )
        targets.append(
            CodingTarget(
                sheet=sheet_name,
                column=column,
                coding_type=coding_type,
                row_count=len(rows),
                non_empty_count=non_empty,
            )
        )

    unique_terms = sorted(unique_map.values(), key=lambda item: (-item.count, item.coding_type, item.term))
    return {
        "filename": filename,
        "sheets": [
            {
                "sheet": sheet_name,
                "headers": headers,
                "row_count": len(rows),
                "coding_columns": [{"column": col, "coding_type": ctype} for col, ctype in sheet_targets],
            }
        ],
        "coding_targets": [t.as_dict() for t in targets],
        "unique_terms": [t.as_dict() for t in unique_terms],
        "stats": {
            "sheet_count": 1,
            "data_rows": len(rows),
            "target_count": len(targets),
            "unique_term_count": len(unique_terms),
        },
    }


def _suggestion_payload(
    term: dict[str, Any],
    match_type: str,
    score: float,
    reason: str,
) -> dict[str, Any]:
    level = term.get("level") or "PT"
    return {
        "level": level,
        "level_label": LEVEL_LABELS.get(level, level),
        "code": term.get("code"),
        "en_name": term.get("en_name") or "",
        "zh_name": term.get("zh_name") or "",
        "is_current": term.get("is_current") or "",
        "match_type": match_type,
        "score": round(float(score), 1),
        "reason": reason,
        "parent_code": term.get("parent_code"),
    }


class TermSuggester:
    """Batch PT/LLT suggestion engine for verbatim coding."""

    def __init__(self, store: Any):
        self._store = store
        self._terms: tuple[dict[str, Any], ...] | None = None
        self._by_level: dict[str, list[dict[str, Any]]] | None = None
        self._norm_index: dict[tuple[str, str], list[dict[str, Any]]] | None = None

    def _ensure_index(self) -> None:
        if self._terms is not None:
            return
        terms = tuple(
            term
            for term in self._store.all_terms()
            if term.get("level") in {"PT", "LLT"}
        )
        by_level: dict[str, list[dict[str, Any]]] = {"PT": [], "LLT": []}
        norm_index: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
        for term in terms:
            level = term.get("level") or ""
            if level in by_level:
                by_level[level].append(term)
            for lang, name in (("zh", term.get("zh_name")), ("en", term.get("en_name"))):
                if not name:
                    continue
                key = (lang, normalize_text(str(name)))
                if key not in norm_index or term not in norm_index[key]:
                    norm_index[key].append(term)
        self._terms = terms
        self._by_level = by_level
        self._norm_index = dict(norm_index)

    def suggest_batch(
        self,
        terms: Iterable[str],
        *,
        mode: str = "both",
        limit_per_term: int = 5,
    ) -> dict[str, dict[str, Any]]:
        self._ensure_index()
        assert self._terms is not None
        assert self._norm_index is not None
        langs = ["zh", "en"] if mode == "both" else [mode]
        limit_per_term = max(1, min(int(limit_per_term), 20))
        results: dict[str, dict[str, Any]] = {}

        for raw_term in terms:
            query = (raw_term or "").strip()
            if not query:
                continue
            results[query] = self._suggest_one(query, langs, limit_per_term)
        return results

    def _suggest_one(self, query: str, langs: list[str], limit_per_term: int) -> dict[str, Any]:
        assert self._norm_index is not None
        assert self._by_level is not None
        assert self._terms is not None

        norm_query = normalize_text(query)
        exact_hits: list[dict[str, Any]] = []
        seen: set[str] = set()

        def key_of(term: dict[str, Any]) -> str:
            return f"{term.get('level')}:{term.get('code')}"

        def push(term: dict[str, Any], match_type: str, score: float, reason: str) -> None:
            key = key_of(term)
            if key in seen:
                return
            seen.add(key)
            exact_hits.append(_suggestion_payload(term, match_type, score, reason))

        for lang in langs:
            for term in self._norm_index.get((lang, norm_query), []):
                # Prefer current LLT for coding; PT is the regulatory aggregation level.
                if term.get("level") == "LLT" and term.get("is_current") != "Y":
                    score = 96
                    reason = "LLT 完全匹配（非当前术语，建议复核）"
                elif term.get("level") == "LLT":
                    score = 100
                    reason = "LLT 完全匹配"
                else:
                    score = 99
                    reason = "PT 完全匹配"
                push(term, "exact", score, reason)

        if not exact_hits:
            # Lexical token-set match (word-order variants).
            query_key = token_key(query)
            if query_key:
                for term in self._terms:
                    for lang in langs:
                        name = term.get("zh_name") if lang == "zh" else term.get("en_name")
                        if name and token_key(str(name)) == query_key:
                            push(term, "lexical", 94, "词元集合一致（词序变体）")
                            break

        if not exact_hits and len(norm_query) >= 2:
            # Containment match: verbatim is inside a MedDRA term or vice versa.
            contains_hits: list[tuple[float, dict[str, Any], str, str]] = []
            for term in self._terms:
                for lang in langs:
                    name = term.get("zh_name") if lang == "zh" else term.get("en_name")
                    if not name:
                        continue
                    norm_name = normalize_text(str(name))
                    if not norm_name or norm_name == norm_query:
                        continue
                    if norm_query in norm_name:
                        # Prefer shorter MedDRA terms (more specific match).
                        score = 84 - min(20, abs(len(norm_name) - len(norm_query)) * 0.3)
                        contains_hits.append((score, term, "contains", "术语文本包含原文"))
                    elif norm_name in norm_query and len(norm_name) >= 2:
                        score = 80 - min(15, abs(len(norm_name) - len(norm_query)) * 0.3)
                        contains_hits.append((score, term, "contains", "原文包含术语文本"))
            contains_hits.sort(
                key=lambda item: (
                    -item[0],
                    0 if item[1].get("level") == "LLT" else 1,
                    abs(len(normalize_text(str(item[1].get("zh_name") or item[1].get("en_name") or ""))) - len(norm_query)),
                )
            )
            for score, term, match_type, reason in contains_hits[:limit_per_term]:
                push(term, match_type, score, reason)

        if not exact_hits and len(norm_query) >= 3:
            fuzzy_hits: list[tuple[float, dict[str, Any], str, str]] = []
            threshold = fuzzy_threshold(norm_query)
            for term in self._terms:
                for lang in langs:
                    name = term.get("zh_name") if lang == "zh" else term.get("en_name")
                    if not name:
                        continue
                    norm_name = normalize_text(str(name))
                    if not fuzzy_candidate(norm_query, norm_name):
                        continue
                    score = fuzzy_score(norm_query, norm_name)
                    if score >= threshold:
                        fuzzy_hits.append((score, term, "fuzzy", "拼写或近似文本候选"))
            fuzzy_hits.sort(key=lambda item: (-item[0], 0 if item[1].get("level") == "LLT" else 1))
            for score, term, match_type, reason in fuzzy_hits[:limit_per_term]:
                push(term, match_type, score, reason)

        # Prefer LLT > PT for coding suggestions at equal score.
        exact_hits.sort(
            key=lambda item: (
                -float(item.get("score") or 0),
                0 if item.get("level") == "LLT" else 1,
                str(item.get("code") or ""),
            )
        )
        best = exact_hits[0] if exact_hits else None
        pt_payload: dict[str, Any] | None = None
        if best and best.get("level") == "LLT" and best.get("parent_code"):
            pt_term = self._find_term("PT", str(best["parent_code"]))
            if pt_term is not None:
                pt_payload = _suggestion_payload(pt_term, "parent", 100, "LLT 对应 PT")
        elif best and best.get("level") == "PT":
            pt_payload = best

        return {
            "query": query,
            "best": best,
            "pt": pt_payload,
            "candidates": exact_hits[:limit_per_term],
            "status": "suggested" if best else "no_match",
        }

    def _find_term(self, level: str, code: str) -> dict[str, Any] | None:
        assert self._by_level is not None
        for term in self._by_level.get(level, []):
            if str(term.get("code")) == str(code):
                return term
        return None


def build_export_rows(
    unique_terms: list[dict[str, Any]],
    decisions: dict[str, dict[str, Any]],
) -> list[dict[str, Any]]:
    """Expand unique-term decisions back into occurrence-level coded listing rows."""
    rows: list[dict[str, Any]] = []
    for item in unique_terms:
        term = item.get("term") or ""
        coding_type = item.get("coding_type") or ""
        decision = decisions.get(f"{coding_type}|{term}") or decisions.get(term) or {}
        suggestion = item.get("suggestion") or {}
        best = suggestion.get("best") or {}
        pt = decision.get("pt") or {}
        llt = decision.get("llt") or {}
        status = decision.get("status") or ("accepted" if best else "pending")
        note = decision.get("note") or ""

        pt_code = decision.get("pt_code") or pt.get("code") or (best.get("code") if best.get("level") == "PT" else "")
        pt_zh = decision.get("pt_zh") or pt.get("zh_name") or (best.get("zh_name") if best.get("level") == "PT" else "")
        pt_en = decision.get("pt_en") or pt.get("en_name") or (best.get("en_name") if best.get("level") == "PT" else "")
        llt_code = decision.get("llt_code") or llt.get("code") or (best.get("code") if best.get("level") == "LLT" else "")
        llt_zh = decision.get("llt_zh") or llt.get("zh_name") or (best.get("zh_name") if best.get("level") == "LLT" else "")
        llt_en = decision.get("llt_en") or llt.get("en_name") or (best.get("en_name") if best.get("level") == "LLT" else "")

        # Fill PT from LLT parent when missing.
        if llt_code and not pt_code:
            pt_code = decision.get("pt_from_llt_code") or ""
            pt_zh = decision.get("pt_from_llt_zh") or ""
            pt_en = decision.get("pt_from_llt_en") or ""

        for occ in item.get("occurrences") or [{}]:
            row = {
                "术语类型": coding_type,
                "原文术语": term,
                "编码状态": status,
                "匹配方式": best.get("match_type") or "",
                "匹配分": best.get("score") or "",
                "PT代码": pt_code,
                "PT中文": pt_zh,
                "PT英文": pt_en,
                "LLT代码": llt_code,
                "LLT中文": llt_zh,
                "LLT英文": llt_en,
                "备注": note,
            }
            for meta in SESSION_META_HEADERS:
                if meta in occ:
                    row[meta] = occ.get(meta) or ""
            if "sheet" in occ:
                row["来源Sheet"] = occ.get("sheet") or ""
            if "column" in occ:
                row["来源列"] = occ.get("column") or ""
            if "excel_row" in occ:
                row["来源行"] = occ.get("excel_row") or ""
            rows.append(row)
    return rows


def code_to_csv(rows: list[dict[str, Any]]) -> str:
    if not rows:
        return ""
    # Stable column order: coding columns first, then metadata.
    preferred = [
        "术语类型",
        "原文术语",
        "编码状态",
        "PT代码",
        "PT中文",
        "PT英文",
        "LLT代码",
        "LLT中文",
        "LLT英文",
        "匹配方式",
        "匹配分",
        "备注",
    ]
    extra_keys: list[str] = []
    for row in rows:
        for key in row.keys():
            if key not in preferred and key not in extra_keys:
                extra_keys.append(key)
    fieldnames = preferred + extra_keys
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=fieldnames, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({key: row.get(key, "") for key in fieldnames})
    return output.getvalue()


def safe_download_filename(name: str, default: str = "meddra_coded_listing.csv") -> str:
    cleaned = re.sub(r'[\\/:*?"<>|\r\n]+', "", Path(name or "").name).strip()
    if not cleaned:
        return default
    if not cleaned.lower().endswith(".csv"):
        cleaned = f"{cleaned}.csv"
    return cleaned[:120]
